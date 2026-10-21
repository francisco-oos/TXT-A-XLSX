"""
TXT -> XLSX | CONVERSOR Y UNIFICADOR
====================================

Menú con 3 modos:

1) INDIVIDUAL
   - Cada TXT se convierte en su propio XLSX.
   - Conserva la misma estructura de carpetas.

2) POR CARPETA / SUBCARPETA
   - Cada carpeta que contenga uno o más TXT genera UN solo XLSX.
   - Cada TXT queda en una hoja diferente dentro de ese XLSX.
   - Se conserva la estructura de carpetas.
   - Incluye una hoja "INDICE" con el origen de cada TXT.

3) TODO EN UN SOLO EXCEL
   - Todos los TXT encontrados dentro de la carpeta raíz y sus subcarpetas
     se guardan en UN solo archivo XLSX.
   - Cada TXT queda en una hoja diferente.
   - Incluye una hoja "INDICE" con la ruta original.

Los TXT originales NO se modifican ni se eliminan.

Requisito:
    pip install openpyxl
"""

from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from pathlib import Path
from tkinter import Tk, StringVar, BooleanVar, filedialog, messagebox
from tkinter import ttk

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment


VERSION = "2.0"

ENCODINGS = (
    "utf-8-sig",
    "utf-8",
    "cp1252",
    "latin-1",
)

DELIMITERS = "\t;,|"
INVALID_SHEET_CHARS_RE = re.compile(r'[\[\]:*?/\\]')
INVALID_FILENAME_CHARS_RE = re.compile(r'[<>:"/\\|?*]')


# ---------------------------------------------------------------------------
# LECTURA Y PARSEO
# ---------------------------------------------------------------------------

def leer_txt(ruta: Path) -> tuple[str, str]:
    """Lee un TXT intentando codificaciones comunes de Windows."""
    ultimo_error = None

    for encoding in ENCODINGS:
        try:
            return ruta.read_text(encoding=encoding), encoding
        except UnicodeDecodeError as exc:
            ultimo_error = exc

    raise UnicodeError(
        f"No fue posible decodificar {ruta}"
    ) from ultimo_error


def detectar_delimitador(texto: str) -> str | None:
    """
    Detecta tabulador, punto y coma, coma o barra vertical.
    Si no parece tabla, devuelve None.
    """
    lineas = [linea for linea in texto.splitlines() if linea.strip()]
    if not lineas:
        return None

    muestra = "\n".join(lineas[:50])

    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=DELIMITERS)
        return dialecto.delimiter
    except csv.Error:
        pass

    mejor = None
    mejor_puntaje = 0

    for delimitador in DELIMITERS:
        conteos = [linea.count(delimitador) for linea in lineas[:50]]
        positivos = [n for n in conteos if n > 0]

        if not positivos:
            continue

        moda = max(set(positivos), key=positivos.count)
        consistentes = sum(1 for n in conteos if n == moda)
        puntaje = consistentes * 100 + moda

        if puntaje > mejor_puntaje:
            mejor_puntaje = puntaje
            mejor = delimitador

    return mejor


def filas_desde_txt(texto: str) -> tuple[list[list[str]], str | None]:
    """Convierte el contenido TXT en filas para Excel."""
    if not texto:
        return [], None

    delimitador = detectar_delimitador(texto)

    if delimitador is None:
        return [[linea] for linea in texto.splitlines()], None

    lector = csv.reader(texto.splitlines(), delimiter=delimitador)
    return [list(fila) for fila in lector], delimitador


def limpiar_valor_excel(valor: str) -> str:
    """Quita caracteres de control no permitidos por XLSX."""
    if valor is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(valor))


def nombre_delimitador(delimitador: str | None) -> str:
    if delimitador == "\t":
        return "TAB"
    if delimitador == ";":
        return ";"
    if delimitador == ",":
        return ","
    if delimitador == "|":
        return "|"
    return "Sin detectar / 1 columna"


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------

def ajustar_anchos(ws, max_ancho: int = 60) -> None:
    """Ajuste visual moderado de columnas."""
    anchos = {}

    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue

            largo = min(len(str(celda.value)) + 2, max_ancho)
            letra = celda.column_letter
            anchos[letra] = max(anchos.get(letra, 0), largo)

    for letra, ancho in anchos.items():
        ws.column_dimensions[letra].width = max(10, ancho)


def nombre_hoja_seguro(nombre: str, usados: set[str]) -> str:
    """Crea un nombre de hoja válido, <=31 caracteres y único."""
    limpio = INVALID_SHEET_CHARS_RE.sub("_", nombre).strip().strip("'")

    if not limpio:
        limpio = "TXT"

    base = limpio[:31]
    candidato = base
    contador = 2

    usados_lower = {x.lower() for x in usados}

    while candidato.lower() in usados_lower:
        sufijo = f"_{contador}"
        candidato = f"{base[:31-len(sufijo)]}{sufijo}"
        contador += 1

    usados.add(candidato)
    return candidato


def nombre_archivo_seguro(nombre: str) -> str:
    limpio = INVALID_FILENAME_CHARS_RE.sub("_", nombre).strip()
    return limpio or "resultado"


def escribir_txt_en_hoja(ws, ruta_txt: Path) -> dict:
    """
    Carga un TXT completo en una worksheet y devuelve metadatos.
    Los valores se guardan como texto para conservar ceros iniciales.
    """
    texto, encoding = leer_txt(ruta_txt)
    filas, delimitador = filas_desde_txt(texto)

    max_columnas = 0

    for numero_fila, fila in enumerate(filas, start=1):
        max_columnas = max(max_columnas, len(fila))

        for numero_columna, valor in enumerate(fila, start=1):
            ws.cell(
                row=numero_fila,
                column=numero_columna,
                value=limpiar_valor_excel(valor),
            )

    if filas:
        ws.freeze_panes = "A2"
        ajustar_anchos(ws)

    return {
        "filas": len(filas),
        "columnas": max_columnas,
        "encoding": encoding,
        "delimitador": nombre_delimitador(delimitador),
    }


def preparar_indice(ws) -> None:
    encabezados = [
        "#",
        "Archivo TXT",
        "Ruta relativa",
        "Hoja",
        "Filas",
        "Columnas",
        "Codificación",
        "Delimitador",
    ]

    ws.append(encabezados)

    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"


def finalizar_indice(ws) -> None:
    ajustar_anchos(ws, max_ancho=80)


def convertir_archivo_individual(origen: Path, destino: Path) -> tuple[int, int]:
    """Convierte un TXT individual en un XLSX."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    meta = escribir_txt_en_hoja(ws, origen)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    wb.close()

    return meta["filas"], meta["columnas"]


def crear_libro_con_varios_txt(
    archivos_txt: list[Path],
    destino: Path,
    origen_raiz: Path,
) -> tuple[int, list[tuple[Path, str]]]:
    """
    Crea UN libro con varios TXT.
    Cada TXT se almacena en una hoja distinta.
    Devuelve cantidad OK y lista de errores.
    """
    wb = Workbook()

    # La hoja inicial será el índice.
    indice = wb.active
    indice.title = "INDICE"
    preparar_indice(indice)

    usados = {"INDICE"}
    correctos = 0
    errores: list[tuple[Path, str]] = []

    for numero, txt in enumerate(archivos_txt, start=1):
        try:
            nombre_hoja = nombre_hoja_seguro(txt.stem, usados)
            ws = wb.create_sheet(title=nombre_hoja)

            meta = escribir_txt_en_hoja(ws, txt)

            relativa = txt.relative_to(origen_raiz)

            indice.append([
                numero,
                txt.name,
                str(relativa),
                nombre_hoja,
                meta["filas"],
                meta["columnas"],
                meta["encoding"],
                meta["delimitador"],
            ])

            correctos += 1

        except Exception as exc:
            errores.append((txt, str(exc)))
            print(f"ERROR en {txt}: {exc}")

    finalizar_indice(indice)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    wb.close()

    return correctos, errores


# ---------------------------------------------------------------------------
# MODOS
# ---------------------------------------------------------------------------

def obtener_txt(origen_raiz: Path) -> list[Path]:
    return sorted(
        (
            p for p in origen_raiz.rglob("*")
            if p.is_file() and p.suffix.lower() == ".txt"
        ),
        key=lambda p: str(p).lower(),
    )


def modo_individual(
    origen_raiz: Path,
    archivos_txt: list[Path],
) -> tuple[Path, int, list[tuple[Path, str]]]:
    """
    Modo 1:
    un XLSX por TXT, conservando la misma estructura.
    """
    destino_raiz = origen_raiz.parent / f"{origen_raiz.name}_XLSX_INDIVIDUAL"

    correctos = 0
    errores = []

    for indice, txt in enumerate(archivos_txt, start=1):
        relativa = txt.relative_to(origen_raiz)
        xlsx = destino_raiz / relativa.with_suffix(".xlsx")

        try:
            filas, columnas = convertir_archivo_individual(txt, xlsx)
            correctos += 1
            print(
                f"[{indice:>4}/{len(archivos_txt)}] OK  "
                f"{relativa} -> {xlsx.relative_to(destino_raiz)} "
                f"({filas} filas, {columnas} columnas)"
            )
        except Exception as exc:
            errores.append((txt, str(exc)))
            print(
                f"[{indice:>4}/{len(archivos_txt)}] ERROR  "
                f"{relativa}: {exc}"
            )

    return destino_raiz, correctos, errores


def modo_por_carpeta(
    origen_raiz: Path,
    archivos_txt: list[Path],
) -> tuple[Path, int, list[tuple[Path, str]]]:
    """
    Modo 2:
    cada carpeta/subcarpeta que tenga TXT genera UN XLSX.
    Cada TXT queda como hoja independiente.
    """
    destino_raiz = origen_raiz.parent / f"{origen_raiz.name}_XLSX_POR_CARPETA"

    grupos: dict[Path, list[Path]] = defaultdict(list)

    # IMPORTANTE:
    # Se agrupan solo los TXT que están directamente dentro de cada carpeta.
    # Un TXT de una subcarpeta pertenece al XLSX de ESA subcarpeta.
    for txt in archivos_txt:
        grupos[txt.parent].append(txt)

    correctos = 0
    errores_totales: list[tuple[Path, str]] = []

    carpetas_ordenadas = sorted(grupos.keys(), key=lambda p: str(p).lower())

    for indice, carpeta in enumerate(carpetas_ordenadas, start=1):
        relativa_carpeta = carpeta.relative_to(origen_raiz)

        if str(relativa_carpeta) == ".":
            nombre_libro = f"{nombre_archivo_seguro(origen_raiz.name)}_RAIZ.xlsx"
            destino = destino_raiz / nombre_libro
            etiqueta = "[RAIZ]"
        else:
            nombre_libro = f"{nombre_archivo_seguro(carpeta.name)}.xlsx"
            destino = destino_raiz / relativa_carpeta / nombre_libro
            etiqueta = str(relativa_carpeta)

        lista = sorted(grupos[carpeta], key=lambda p: p.name.lower())

        ok, errores = crear_libro_con_varios_txt(
            lista,
            destino,
            origen_raiz,
        )

        correctos += ok
        errores_totales.extend(errores)

        print(
            f"[{indice:>4}/{len(carpetas_ordenadas)}] "
            f"CARPETA {etiqueta} -> {destino.relative_to(destino_raiz)} "
            f"({ok}/{len(lista)} TXT)"
        )

    return destino_raiz, correctos, errores_totales


def modo_unificado(
    origen_raiz: Path,
    archivos_txt: list[Path],
) -> tuple[Path, int, list[tuple[Path, str]]]:
    """
    Modo 3:
    todos los TXT de toda la estructura en UN XLSX.
    """
    destino_raiz = origen_raiz.parent / f"{origen_raiz.name}_XLSX_UNIFICADO"

    nombre = f"{nombre_archivo_seguro(origen_raiz.name)}_TODO.xlsx"
    destino = destino_raiz / nombre

    correctos, errores = crear_libro_con_varios_txt(
        archivos_txt,
        destino,
        origen_raiz,
    )

    print(
        f"UNIFICADO -> {destino} "
        f"({correctos}/{len(archivos_txt)} TXT)"
    )

    return destino_raiz, correctos, errores


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

class Aplicacion:
    def __init__(self, root: Tk):
        self.root = root
        self.root.title(f"TXT -> XLSX | Conversor y Unificador v{VERSION}")
        self.root.resizable(False, False)

        self.carpeta = StringVar(value="")
        self.modo = StringVar(value="1")
        self.en_ejecucion = BooleanVar(value=False)

        self._crear_interfaz()

    def _crear_interfaz(self) -> None:
        marco = ttk.Frame(self.root, padding=18)
        marco.grid(row=0, column=0, sticky="nsew")

        titulo = ttk.Label(
            marco,
            text="TXT → XLSX",
            font=("Segoe UI", 18, "bold"),
        )
        titulo.grid(row=0, column=0, columnspan=3, sticky="w")

        subtitulo = ttk.Label(
            marco,
            text="Convierte, agrupa por carpetas o unifica todos tus TXT.",
        )
        subtitulo.grid(row=1, column=0, columnspan=3, sticky="w", pady=(0, 16))

        ttk.Label(
            marco,
            text="Carpeta origen:",
            font=("Segoe UI", 10, "bold"),
        ).grid(row=2, column=0, columnspan=3, sticky="w")

        entrada = ttk.Entry(
            marco,
            textvariable=self.carpeta,
            width=70,
            state="readonly",
        )
        entrada.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(4, 12))

        ttk.Button(
            marco,
            text="Seleccionar carpeta…",
            command=self.seleccionar_carpeta,
        ).grid(row=3, column=2, padx=(8, 0), pady=(4, 12))

        ttk.Separator(marco).grid(
            row=4, column=0, columnspan=3, sticky="ew", pady=(2, 12)
        )

        ttk.Label(
            marco,
            text="Elige qué quieres hacer:",
            font=("Segoe UI", 10, "bold"),
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(0, 6))

        opciones = [
            (
                "1",
                "1. Un XLSX por cada TXT",
                "Conserva la misma estructura de carpetas. "
                "Es el comportamiento de la versión anterior.",
            ),
            (
                "2",
                "2. Un XLSX por cada carpeta/subcarpeta",
                "Si una subcarpeta contiene 2, 5 o 20 TXT, "
                "los mete en un solo Excel; cada TXT queda como una hoja.",
            ),
            (
                "3",
                "3. Todos los TXT en un solo XLSX",
                "Busca en toda la carpeta raíz y sus subcarpetas. "
                "Cada TXT queda como una hoja y se agrega un índice.",
            ),
        ]

        fila = 6
        for valor, texto, descripcion in opciones:
            ttk.Radiobutton(
                marco,
                text=texto,
                value=valor,
                variable=self.modo,
            ).grid(row=fila, column=0, columnspan=3, sticky="w")

            ttk.Label(
                marco,
                text=descripcion,
                foreground="#555555",
                wraplength=650,
            ).grid(
                row=fila + 1,
                column=0,
                columnspan=3,
                sticky="w",
                padx=(24, 0),
                pady=(0, 10),
            )
            fila += 2

        ttk.Separator(marco).grid(
            row=fila, column=0, columnspan=3, sticky="ew", pady=(4, 12)
        )

        ttk.Label(
            marco,
            text="Los TXT originales nunca se modifican.",
        ).grid(row=fila + 1, column=0, columnspan=2, sticky="w")

        self.boton_ejecutar = ttk.Button(
            marco,
            text="PROCESAR",
            command=self.procesar,
        )
        self.boton_ejecutar.grid(
            row=fila + 1,
            column=2,
            sticky="e",
            ipadx=18,
            ipady=6,
        )

    def seleccionar_carpeta(self) -> None:
        carpeta = filedialog.askdirectory(
            title="Selecciona la carpeta raíz que contiene los TXT"
        )
        if carpeta:
            self.carpeta.set(carpeta)

    def procesar(self) -> None:
        if self.en_ejecucion.get():
            return

        carpeta = self.carpeta.get().strip()

        if not carpeta:
            messagebox.showwarning(
                "Falta carpeta",
                "Primero selecciona la carpeta que contiene los TXT.",
            )
            return

        origen_raiz = Path(carpeta).resolve()

        if not origen_raiz.is_dir():
            messagebox.showerror(
                "Carpeta no válida",
                f"La carpeta ya no existe:\n\n{origen_raiz}",
            )
            return

        archivos_txt = obtener_txt(origen_raiz)

        if not archivos_txt:
            messagebox.showwarning(
                "Sin archivos TXT",
                f"No se encontraron archivos .txt dentro de:\n\n{origen_raiz}",
            )
            return

        self.en_ejecucion.set(True)
        self.boton_ejecutar.configure(state="disabled")
        self.root.update_idletasks()

        print("=" * 90)
        print(f"TXT -> XLSX | v{VERSION}")
        print("=" * 90)
        print(f"Origen: {origen_raiz}")
        print(f"TXT encontrados: {len(archivos_txt)}")
        print(f"Modo seleccionado: {self.modo.get()}")
        print("=" * 90)

        try:
            if self.modo.get() == "1":
                destino, correctos, errores = modo_individual(
                    origen_raiz,
                    archivos_txt,
                )
                nombre_modo = "Un XLSX por cada TXT"

            elif self.modo.get() == "2":
                destino, correctos, errores = modo_por_carpeta(
                    origen_raiz,
                    archivos_txt,
                )
                nombre_modo = "Un XLSX por carpeta/subcarpeta"

            else:
                destino, correctos, errores = modo_unificado(
                    origen_raiz,
                    archivos_txt,
                )
                nombre_modo = "Todos los TXT en un solo XLSX"

            print("=" * 90)
            print(f"Correctos: {correctos}")
            print(f"Errores: {len(errores)}")
            print(f"Resultado: {destino}")
            print("=" * 90)

            mensaje = (
                f"Proceso terminado.\n\n"
                f"Modo: {nombre_modo}\n"
                f"TXT encontrados: {len(archivos_txt)}\n"
                f"Procesados correctamente: {correctos}\n"
                f"Errores: {len(errores)}\n\n"
                f"Los TXT originales NO fueron modificados.\n\n"
                f"Resultado:\n{destino}"
            )

            if errores:
                detalle = "\n".join(
                    f"- {ruta.name}: {error}"
                    for ruta, error in errores[:10]
                )
                mensaje += f"\n\nPrimeros errores:\n{detalle}"

            messagebox.showinfo("TXT -> XLSX", mensaje)

        except Exception as exc:
            messagebox.showerror(
                "Error inesperado",
                f"Ocurrió un error:\n\n{exc}",
            )
            print(f"ERROR INESPERADO: {exc}")

        finally:
            self.en_ejecucion.set(False)
            self.boton_ejecutar.configure(state="normal")


def main() -> int:
    root = Tk()

    # Intenta mantener la ventana al frente al iniciar.
    try:
        root.attributes("-topmost", True)
        root.after(600, lambda: root.attributes("-topmost", False))
    except Exception:
        pass

    Aplicacion(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
